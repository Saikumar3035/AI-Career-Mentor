import json
import logging
import os
import re
import tempfile
from typing import Dict, List

import faiss
import streamlit as st
from dotenv import load_dotenv
from docx import Document
from openai import OpenAI
from pdfminer.high_level import extract_text
from sentence_transformers import SentenceTransformer

from rag_star.config import (
    CAREER_PATHS,
    CERTIFICATE_MAP,
    KNOWLEDGE_BASE_FILE,
    PROJECT_TEMPLATES,
    ROLE_PATTERNS,
    SKILL_DATABASE,
    SOFT_SKILLS,
)

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)
load_dotenv()


class CareerMentorEngine:
    def __init__(self):
        self.openai_client = None
        self.knowledge_base = self.load_knowledge_base()
        self.sentence_model = self.get_sentence_model()
        self._init_openai_client()

    def _init_openai_client(self):
        api_key = os.getenv("OPENAI_API_KEY")
        if api_key:
            try:
                self.openai_client = OpenAI(api_key=api_key)
            except Exception as exc:
                logger.warning(f"OpenAI failed to initialize: {exc}")
                self.openai_client = None

    def get_sentence_model(self):
        if "sentence_model" not in st.session_state:
            with st.spinner("Loading AI embedding model..."):
                st.session_state.sentence_model = SentenceTransformer("all-MiniLM-L6-v2")
        return st.session_state.sentence_model

    def load_knowledge_base(self):
        if not KNOWLEDGE_BASE_FILE.exists():
            data = [
                {"text": "Python is one of the most practical languages for AI projects, data engineering, automation, and backend services.", "meta": {"skill": "Python", "platform": "Coursera", "duration": "40 hours", "level": "Beginner"}},
                {"text": "Machine learning requires feature engineering, evaluation metrics, model comparison, and deployment thinking.", "meta": {"skill": "Machine Learning", "platform": "Coursera", "duration": "55 hours", "level": "Intermediate"}},
                {"text": "SQL is essential for analytics, data pipelines, business dashboards, and AI product data queries.", "meta": {"skill": "SQL", "platform": "SQLBolt", "duration": "20 hours", "level": "Beginner"}},
                {"text": "Docker helps developers build portable and repeatable environments for MLOps and deployments.", "meta": {"skill": "Docker", "platform": "Udemy", "duration": "30 hours", "level": "Intermediate"}},
                {"text": "AWS architecture for AI and software projects includes IAM, EC2, S3, Lambda, and secure deployment workflows.", "meta": {"skill": "AWS", "platform": "A Cloud Guru", "duration": "45 hours", "level": "Intermediate"}},
            ]
            KNOWLEDGE_BASE_FILE.parent.mkdir(parents=True, exist_ok=True)
            KNOWLEDGE_BASE_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")

        with open(KNOWLEDGE_BASE_FILE, "r", encoding="utf-8") as file:
            return json.load(file)

    def extract_text_from_file(self, file, file_type: str) -> str:
        try:
            if file_type.lower() == "pdf":
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as temp:
                    temp.write(file.read())
                    path = temp.name
                try:
                    return extract_text(path) or ""
                finally:
                    os.unlink(path)

            if file_type.lower() == "docx":
                with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as temp:
                    temp.write(file.read())
                    path = temp.name
                try:
                    doc = Document(path)
                    paragraphs = [p.text for p in doc.paragraphs]
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                paragraphs.append(cell.text)
                    return "\n".join(paragraphs)
                finally:
                    os.unlink(path)

            return file.read().decode("utf-8", errors="ignore")
        except Exception as exc:
            logger.error(f"Text extraction failed: {exc}")
            return ""

    def extract_technical_skills(self, text: str) -> List[str]:
        found = []
        lower = text.lower()
        for skills in SKILL_DATABASE.values():
            for skill in skills:
                if re.search(r"\b" + re.escape(skill.lower()) + r"\b", lower):
                    found.append(skill)
        return sorted(set(found))

    def extract_soft_skills(self, text: str) -> List[str]:
        lower = text.lower()
        return [skill for skill in SOFT_SKILLS if skill.lower() in lower]

    def extract_experience_years(self, text: str) -> int:
        patterns = [
            r"(\d+)\+?\s*years?\s*(?:of\s*)?experience",
            r"experience.*?(\d+)\+?\s*years?",
            r"(\d+)\+?\s*yrs?\s*(?:of\s*)?experience",
        ]
        values = []
        for pattern in patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            for match in matches:
                try:
                    values.append(int(match))
                except ValueError:
                    pass
        return max(values) if values else 0

    def extract_education(self, text: str) -> List[str]:
        patterns = [
            r"(?:Bachelor|Master|PhD|B\.S\.|M\.S\.|B\.Tech|M\.Tech|B\.A\.|M\.A\.).*?(?:in|of)\s*([^,\n]+)",
            r"(?:University|College|Institute).*?([^,\n]+)",
            r"(Computer Science|Engineering|Mathematics|Physics|Statistics|Business)",
        ]
        results = []
        for pattern in patterns:
            results.extend(re.findall(pattern, text, re.IGNORECASE))
        return list(dict.fromkeys(results))[:3]

    def extract_certifications(self, text: str) -> List[str]:
        results = []
        for keyword in CERTIFICATE_MAP:
            if re.search(re.escape(keyword), text, re.IGNORECASE):
                results.extend(CERTIFICATE_MAP[keyword])
        return list(dict.fromkeys(results))[:5]

    def extract_projects(self, text: str) -> List[str]:
        patterns = [
            r"(?:Project|Built|Developed|Created|Implemented|Designed).*?([^.!?\n]+)",
            r"(?:GitHub|Portfolio).*?([^.!?\n]+)",
        ]
        results = []
        for pattern in patterns:
            results.extend(match.strip() for match in re.findall(pattern, text, re.IGNORECASE))
        return list(dict.fromkeys(results))[:5]

    def extract_companies(self, text: str) -> List[str]:
        patterns = [
            r"(?:at|@)\s+([A-Z][A-Za-z0-9&.\s-]+?)(?:\s|,|\n|\.)",
            r"(?:Company|Corp|Inc|Ltd|LLC|Technologies|Systems|Solutions)\s*([^,\n]+)",
        ]
        results = []
        for pattern in patterns:
            results.extend(re.findall(pattern, text))
        return list(dict.fromkeys(results))[:5]

    def extract_roles(self, text: str) -> List[str]:
        results = []
        for pattern in ROLE_PATTERNS:
            results.extend(re.findall(pattern, text, re.IGNORECASE))
        return list(dict.fromkeys(results))[:5]

    def extract_comprehensive_info(self, text: str, is_resume: bool = True) -> Dict:
        info = {
            "skills": {
                "technical": self.extract_technical_skills(text),
                "soft": self.extract_soft_skills(text),
                "all": [],
            },
            "experience_years": 0,
            "education": [],
            "certifications": [],
            "projects": [],
            "companies": [],
            "roles": [],
        }
        info["skills"]["all"] = info["skills"]["technical"] + info["skills"]["soft"]

        if is_resume:
            info["experience_years"] = self.extract_experience_years(text)
            info["education"] = self.extract_education(text)
            info["certifications"] = self.extract_certifications(text)
            info["projects"] = self.extract_projects(text)
            info["companies"] = self.extract_companies(text)
            info["roles"] = self.extract_roles(text)
        return info

    def calculate_match_score(self, resume_skills: List[str], job_skills: List[str]) -> Dict:
        resume_lower = {skill.lower() for skill in resume_skills}
        matched = [skill for skill in job_skills if skill.lower() in resume_lower]
        missing = [skill for skill in job_skills if skill.lower() not in resume_lower]
        total_required = len(job_skills)
        match_pct = (len(matched) / total_required * 100) if total_required else 0

        return {
            "match_percentage": round(match_pct, 1),
            "matched_skills": matched,
            "missing_skills": missing,
            "total_resume_skills": len(resume_skills),
            "total_job_skills": total_required,
            "skill_gap": len(missing),
        }

    def suggest_career_path(self, skills: List[str]) -> Dict:
        results = {}
        skill_set = {skill.lower() for skill in skills}
        for path_name, path_info in CAREER_PATHS.items():
            all_skills = path_info["core_skills"] + path_info["advanced_skills"] + path_info["tools"]
            matched = len(skill_set.intersection({skill.lower() for skill in all_skills}))
            score = (matched / len(all_skills) * 100) if all_skills else 0
            results[path_name] = {
                "score": round(score, 1),
                "matching_skills": matched,
                "missing_core": [skill for skill in path_info["core_skills"] if skill.lower() not in skill_set],
            }
        return dict(sorted(results.items(), key=lambda item: item[1]["score"], reverse=True))

    def get_learning_resources(self, query: str, top_k: int = 5) -> List[Dict]:
        if not self.knowledge_base or not self.sentence_model:
            return []
        try:
            texts = [item["text"] for item in self.knowledge_base]
            embeddings = self.sentence_model.encode(texts, convert_to_numpy=True).astype("float32")
            index = faiss.IndexFlatL2(embeddings.shape[1])
            index.add(embeddings)
            query_embedding = self.sentence_model.encode([query], convert_to_numpy=True).astype("float32")
            distances, indices = index.search(query_embedding, k=min(top_k, len(self.knowledge_base)))
            results = []
            for dist, idx in zip(distances[0], indices[0]):
                if idx < 0:
                    continue
                item = self.knowledge_base[int(idx)].copy()
                item["similarity_score"] = float(1 / (1 + dist))
                results.append(item)
            return results
        except Exception as exc:
            logger.error(f"RAG resource lookup failed: {exc}")
            return []

    def generate_project_ideas(self, missing_skills: List[str]) -> List[str]:
        ideas = []
        for skill in missing_skills[:3]:
            if skill in PROJECT_TEMPLATES:
                ideas.append(f"{skill}: {PROJECT_TEMPLATES[skill]}")
        return ideas

    def suggest_certifications(self, missing_skills: List[str]) -> List[str]:
        suggestions = []
        for skill in missing_skills:
            for key, certs in CERTIFICATE_MAP.items():
                if key.lower() in skill.lower():
                    suggestions.extend(certs)
        return list(dict.fromkeys(suggestions))[:3]

    def get_personalized_recommendations(self, resume_info: Dict, job_info: Dict, match_score: Dict) -> Dict:
        missing = match_score["missing_skills"]
        return {
            "priority_skills": missing[:5],
            "learning_resources": self.get_learning_resources(" ".join(missing), top_k=3) if missing else [],
            "project_ideas": self.generate_project_ideas(missing),
            "certification_suggestions": self.suggest_certifications(missing),
            "resume_improvements": [
                "Quantify your impact with measurable outcomes.",
                "Tailor your skills section to the target role.",
                "Highlight projects that show job-relevant work."
            ],
        }

    def generate_fallback_insights(self, resume_info: Dict, job_info: Dict, match_score: Dict) -> str:
        percent = match_score["match_percentage"]
        if percent >= 80:
            readiness = "You are close to strong competitive readiness."
        elif percent >= 60:
            readiness = "You have a solid base and just need a few focused gaps."
        else:
            readiness = "You need a tighter skill plan before applying at scale."
        priority = match_score["missing_skills"][0] if match_score["missing_skills"] else "portfolio building"
        timeline = "2-3 months" if percent >= 60 else "3-6 months"
        return (
            f"**Readiness:** {readiness}\n"
            f"**Top priority:** strengthen {priority} with a project and practice.\n"
            f"**Timeline:** {timeline} with consistent effort.\n"
            f"**Confidence:** Your current foundation is already promising."
        )

    def generate_ai_insights(self, resume_info: Dict, job_info: Dict, match_score: Dict) -> str:
        if not self.openai_client:
            return self.generate_fallback_insights(resume_info, job_info, match_score)

        prompt = f"""
        Resume skills: {resume_info['skills']['all']}
        Job skills: {job_info['skills']['all']}
        Match percentage: {match_score['match_percentage']}%
        Missing skills: {match_score['missing_skills']}

        Give 4 short sections: Readiness, Priority Actions, Timeline, Confidence Builder.
        Keep the answer concise and motivating.
        """

        try:
            response = self.openai_client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[{"role": "user", "content": prompt}],
                max_tokens=300,
                temperature=0.7,
            )
            return response.choices[0].message.content
        except Exception as exc:
            logger.warning(f"OpenAI insight generation failed: {exc}")
            return self.generate_fallback_insights(resume_info, job_info, match_score)
